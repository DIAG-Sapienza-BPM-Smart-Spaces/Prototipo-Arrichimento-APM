#!/usr/bin/env python3
"""
name_eval.py — score the induced taxonomy against a small gold set.

WHY THIS EXISTS
---------------
Every change to regime C so far has been judged by eye. That failed:
feeding real records to the naming step RAISED slice-to-slice agreement
(6 -> 7 shared categories) while LOWERING quality (freni -> PNEUMATICI,
riscaldamento -> RAFFREDDAMENTO). Stability is gameable by being
consistently wrong. Only correctness against a fixed reference is not.

WHAT IS SCORED
--------------
One row per distinct phrase. Gold = the category that phrase SHOULD end
up in. Prediction = the category it DID end up in, read from
induced_taxonomy__*.json. This scores clustering and naming together,
end to end, which is what we actually care about.

Accuracy is weighted by how many records carry the phrase, so a mistake
on 'clima' (33 records) costs more than one on 'lavavetri' (2).

NOTE ON THE SHEET
-----------------
The sheet carries the real maintenance records behind each phrase, because
a phrase alone ("sedile autista") is not enough to label reliably. The
records misled the 8B model — they do not mislead a human reading them.
By default only the top N phrases are exported: the long tail is 1-2
records each and is not worth anyone's time.

USAGE
-----
  python3 name_eval.py --make-sheet            # top 25 phrases
  python3 name_eval.py --make-sheet --top 0    # all of them
  # fill the `gold` column in name_gold.xlsx, then:
  python3 name_eval.py --score induced_taxonomy__0_150.json \
                          --score induced_taxonomy__150_150.json
  python3 name_eval.py --self-test
"""
import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
import config
DEFAULT_TOP = 25
EXAMPLES_IN_SHEET = 4
EXAMPLE_CHARS = 160


def _default_dir():
    """The directory taxonomy_bootstrap.py writes to.

    Resolved the same way taxonomy_bootstrap does (config.TAXONOMY_FREE_DIR)
    so this script works from any working directory, with fallbacks if it is
    run outside the project.
    """
    try:
        import config
        d = config.TAXONOMY_FREE_DIR
        if d.is_dir():
            return d
    except Exception:
        pass
    for cand in (HERE / "taxonomy_free", HERE, Path.cwd() / "taxonomy_free",
                 Path.cwd()):
        if cand.is_dir() and any(cand.glob("c_descriptions__*.json")):
            return cand
    return HERE


OUT_DIR = _default_dir()
SHEET_NAME = "name_gold.xlsx"

NUOVO = "NUOVO"        # prefix: no reference term fits, a NEW category is correct
SKIP = "SKIP"          # undecidable, leave out of scoring
ALT = "|"              # separator for equally acceptable answers


def canon(name):
    """Put a category name in the same form the pipeline would produce.

    taxonomy_bootstrap.taxonomy_key() strips the IMPIANTO_/SISTEMA_ prefix, so
    the system can never emit "IMPIANTO_IDRAULICO" — only "IDRAULICO".
    Writing the full form in the gold column would score a correct answer
    as wrong. Both sides are normalised through the pipeline's own rule so
    the comparison measures meaning, not spelling convention.
    """
    try:
        from taxonomy_bootstrap import taxonomy_key
        return taxonomy_key(name)
    except Exception:
        import re
        s = re.sub(r"[^A-ZÀ-Ù0-9\s_]", " ", str(name).upper().strip())
        s = re.sub(r"\s+", "_", s.strip())
        for pre in ("SISTEMA_DI_", "SISTEMA_DEL_", "SISTEMA_DELLA_", "SISTEMA_",
                    "IMPIANTO_DI_", "IMPIANTO_", "CIRCUITO_DI_", "CIRCUITO_",
                    "GRUPPO_"):
            if s.startswith(pre):
                return s[len(pre):] or s
        return s


def load_standard_terms():
    """Names the system can pick off the shelf, i.e. NOT inventions.

    NUOVO is only satisfied if the system produced something outside this
    set. Otherwise 'invent a new category' would be scored as passed
    whenever the system quietly force-fitted an existing term — the exact
    blindness this gold set exists to expose.
    """
    try:
        from name_normalizer import load_reference
        return {t.strip().upper() for t in load_reference()}
    except Exception:
        return set()


# --------------------------------------------------------------------------
# reading what the pipeline produced
# --------------------------------------------------------------------------
def _stripper():
    try:
        from taxonomy_bootstrap import strip_boilerplate
        return strip_boilerplate
    except Exception:
        import re

        def strip_boilerplate(p):
            p = str(p).lower().strip()
            p = re.sub(r"^(impianto|sistema)\s+(di\s+|del\s+|dell'|d')?", "", p)
            return re.sub(r"[^a-z0-9\s]", " ", p).strip()
        return strip_boilerplate


def load_descriptions(paths):
    """Cached phase-1 files -> (counts, examples).

    counts:   phrase -> number of records
    examples: phrase -> list of the original record texts
    """
    strip = _stripper()
    counts = Counter()
    examples = defaultdict(list)
    for p in paths:
        data = json.loads(Path(p).read_text(encoding="utf-8"))
        rows = data["pairs"] if isinstance(data, dict) and "pairs" in data else data
        for row in rows:
            if isinstance(row, dict):
                text = row.get("text") or row.get("testo") or ""
                phrase = row.get("phrase") or row.get("descrizione")
            else:
                text, phrase = row[0], row[1]
            key = strip(phrase)
            counts[key] += 1
            if text:
                examples[key].append(str(text))
    return counts, examples


def phrase_to_category(taxonomy_path):
    """induced_taxonomy__*.json -> {phrase: category key it landed in}."""
    data = json.loads(Path(taxonomy_path).read_text(encoding="utf-8"))
    mapping = {}
    for cat in data.get("_provenance", []):
        for phrase in cat["phrases"]:
            mapping[phrase] = cat["key"]
    return mapping


# --------------------------------------------------------------------------
# sheet
# --------------------------------------------------------------------------
def _clip(t):
    t = " ".join(str(t).split())
    return t[:EXAMPLE_CHARS] + ("..." if len(t) > EXAMPLE_CHARS else "")


def make_sheet(counts, examples, hints, out, top=DEFAULT_TOP):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    ordered = counts.most_common(top if top else None)

    wb = Workbook()
    ws = wb.active
    ws.title = "gold"
    ws.append(["phrase", "records", "gold", "real maintenance records",
               "hint (candidates)", "note"])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for phrase, n in ordered:
        ex = "\n".join(f"• {_clip(t)}"
                       for t in examples.get(phrase, [])[:EXAMPLES_IN_SHEET])
        ws.append([phrase, n, "", ex, ", ".join(hints.get(phrase, [])), ""])

    for col, w in {"A": 28, "B": 8, "C": 24, "D": 78, "E": 40, "F": 28}.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row[0].row].height = 62
    ws.freeze_panes = "C2"

    note = wb.create_sheet("istruzioni")
    for line in [
        ["Fill column C (`gold`): which category SHOULD this phrase belong to?"],
        [""],
        ["Column D shows the real maintenance records behind the phrase."],
        ["Decide from those, not from the phrase alone."],
        [""],
        ["IMPORTANT: do not just copy the first hint in column E."],
        ["The hints come from the embedding model, which is part of the system"],
        ["being scored. Copying them measures the system against itself."],
        [""],
        [f"Write {NUOVO} (or {NUOVO}_QUALCOSA) if no existing term fits and the"],
        ["system SHOULD invent a new category. The suffix is only a note for you;"],
        ["scoring just checks that the system invented instead of reusing a term."],
        [f"Write A{ALT}B if two answers are equally acceptable, e.g. SOSPENSIONI{ALT}MOTORE."],
        [f"Write {SKIP} (or leave blank) if you cannot decide. Do not guess —"],
        ["a small clean gold set beats a large noisy one."],
        [""],
        ["Rows are sorted by record count, so the top rows matter most."],
        ["Put reasoning for the debatable ones in column F; those are the"],
        ["genuine domain questions and are worth showing your professor."],
    ]:
        note.append(line)
    note.column_dimensions["A"].width = 84

    wb.save(out)
    return out, len(ordered)


def read_gold(path):
    from openpyxl import load_workbook
    ws = load_workbook(path)["gold"]
    gold = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        phrase, n, g = row[0], row[1], row[2]
        if phrase is None:
            continue
        g = str(g or "").strip().upper()
        if g and g != SKIP:
            gold[str(phrase)] = (g, int(n or 0))
    return gold


# --------------------------------------------------------------------------
# scoring
# --------------------------------------------------------------------------
def judge(want, got, standard):
    """Is `got` an acceptable answer for gold value `want`?

    want may be:
      FRENI              exact category
      FRENI|PNEUMATICI   either one is acceptable (genuine ambiguity)
      NUOVO / NUOVO_X    any category the system INVENTED, i.e. one that is
                         not in the standard reference vocabulary. The
                         suffix after NUOVO is documentation for the human
                         and is deliberately not string-matched.
    """
    got_c = canon(got)
    for raw in want.split(ALT):
        if not raw.strip():
            continue
        if raw.strip().upper().startswith(NUOVO):
            if got_c not in {canon(t) for t in standard}:
                return True
        elif got_c == canon(raw):
            return True
    return False


def score(gold, pred, standard=None):
    standard = load_standard_terms() if standard is None else standard
    rows, hit_w, tot_w, hit_n, tot_n = [], 0, 0, 0, 0
    for phrase, (want, n) in sorted(gold.items(), key=lambda kv: -kv[1][1]):
        got = pred.get(phrase)
        if got is None:
            outcome = "uncovered"      # below min_support, never categorised
        else:
            outcome = "ok" if judge(want, got, standard) else "wrong"
        if outcome != "uncovered":
            tot_w += n
            tot_n += 1
            if outcome == "ok":
                hit_w += n
                hit_n += 1
        rows.append((phrase, n, want, got or "-", outcome))
    return rows, {
        "weighted": hit_w / tot_w if tot_w else 0.0,
        "unweighted": hit_n / tot_n if tot_n else 0.0,
        "covered": tot_n,
        "uncovered": sum(1 for r in rows if r[4] == "uncovered"),
    }


def report(name, rows, stats):
    print(f"\n=== {name} ===")
    print(f"  weighted accuracy   {stats['weighted']:.1%}  "
          "(by records — big phrases matter more)")
    print(f"  unweighted accuracy {stats['unweighted']:.1%}  (by phrase)")
    print(f"  scored {stats['covered']} phrases, "
          f"{stats['uncovered']} uncovered (below min support)")
    bad = [r for r in rows if r[4] == "wrong"]
    if bad:
        print(f"\n  {len(bad)} wrong:")
        for phrase, n, want, got, _ in bad:
            print(f"    {phrase:30} {n:3}rec  want {want:24} got {got}")


# --------------------------------------------------------------------------
def self_test():
    gold = {"freni": ("FRENI", 11), "clima": ("CLIMATIZZAZIONE", 20),
            "idraulico": ("IDRAULICO", 8), "rara": ("TRASMISSIONE", 1)}
    pred = {"freni": "PNEUMATICI", "clima": "CLIMATIZZAZIONE",
            "idraulico": "IDRAULICO"}
    rows, stats = score(gold, pred, set())
    assert stats["uncovered"] == 1 and stats["covered"] == 3, stats
    assert abs(stats["weighted"] - 28 / 39) < 1e-9, stats["weighted"]
    assert abs(stats["unweighted"] - 2 / 3) < 1e-9, stats["unweighted"]
    assert {r[0]: r[4] for r in rows} == {
        "freni": "wrong", "clima": "ok",
        "idraulico": "ok", "rara": "uncovered"}

    # a consistently-wrong system scores 0 even though it is perfectly
    # stable — the exact failure mode this tool exists to catch.
    assert score({"freni": ("FRENI", 11)}, {"freni": "PNEUMATICI"}, set())[1]["weighted"] == 0.0

    std = {"FRENI", "MOTORE", "SOSPENSIONI"}
    # alternatives: either answer accepted
    assert judge("SOSPENSIONI|MOTORE", "MOTORE", std)
    assert judge("SOSPENSIONI|MOTORE", "SOSPENSIONI", std)
    assert not judge("SOSPENSIONI|MOTORE", "FRENI", std)
    # NUOVO demands a genuine invention, not a shelf term
    assert judge("NUOVO", "CENTRALINE_ELETTRONICHE", std)
    assert not judge("NUOVO", "MOTORE", std)
    # the suffix is documentation only, it is never string-matched
    assert judge("NUOVO_ELETTRONICO", "QUALSIASI_COSA_NUOVA", std)
    assert not judge("NUOVO_ELETTRONICO", "FRENI", std)
    # case and spacing tolerated
    assert judge(" freni | motore ", "FRENI", std)
    # the IMPIANTO_ prefix is a spelling convention, not a difference
    assert judge("IMPIANTO_IDRAULICO", "IDRAULICO", std)
    assert judge("IDRAULICO", "IMPIANTO_IDRAULICO", std)
    assert judge("IMPIANTO_ELETTRICO", "ELETTRICO", std)
    assert not judge("IMPIANTO_ELETTRICO", "CENTRALINE_ELETTRONICHE", std)
    # NUOVO still rejects a shelf term written in either form
    assert not judge("NUOVO", "IMPIANTO_FRENI", {"FRENI"})

    # sheet respects --top and carries examples
    counts = Counter({"a": 5, "b": 3, "c": 1})
    ex = {"a": ["record one", "record two"], "b": ["record three"]}
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        out, n = make_sheet(counts, ex, {}, Path(d) / "s.xlsx", top=2)
        assert n == 2, n
        from openpyxl import load_workbook
        ws = load_workbook(out)["gold"]
        assert ws.max_row == 3, ws.max_row
        assert "record one" in ws.cell(2, 4).value
    print("self-test OK")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--make-sheet", action="store_true")
    ap.add_argument("--top", type=int, default=DEFAULT_TOP,
                    help=f"how many phrases to export (0 = all, default {DEFAULT_TOP})")
    ap.add_argument("--sheet", default=None,
                    help=f"gold workbook (default: {config.EVAL_DIR / SHEET_NAME})")
    ap.add_argument("--dir", default=None,
                    help=f"folder holding the run outputs (default: {OUT_DIR})")
    ap.add_argument("--desc", nargs="*", default=None)
    ap.add_argument("--score", action="append", default=[])
    ap.add_argument("--self-test", action="store_true")
    a = ap.parse_args()

    if a.self_test:
        self_test()
        return

    work = Path(a.dir).expanduser().resolve() if a.dir else OUT_DIR
    if a.sheet:
        sheet = Path(a.sheet).expanduser()
        if not sheet.is_absolute() and not sheet.exists():
            sheet = work / a.sheet
    else:
        sheet = config.EVAL_DIR / SHEET_NAME

    if a.make_sheet:
        paths = [Path(p) for p in a.desc] if a.desc \
            else sorted(work.glob("c_descriptions__*.json"))
        if not paths:
            sys.exit(f"No c_descriptions__*.json in {work}\n"
                     f"Use --dir, or run taxonomy_bootstrap.py first.")
        print(f"Folder: {work}")
        for p in paths:
            print(f"  reading {p.name}")
        counts, examples = load_descriptions(paths)
        if not any(examples.values()):
            print("  WARNING: cached files carry no record text; "
                  "the examples column will be empty.")

        keep = [p for p, _ in counts.most_common(a.top if a.top else None)]
        hints = {}
        try:
            from name_normalizer import load_reference, shortlist
            from validator import embedder
            reference = load_reference()
            ref_vecs = {t: embedder.embed_query(f"{t}: {d}")
                        for t, d in reference.items()}
            print(f"Adding candidate hints for {len(keep)} phrases...")
            for phrase in keep:
                v = embedder.embed_query(phrase)
                hints[phrase] = [t for t, _ in shortlist(v, ref_vecs, 4)]
        except Exception as e:
            print(f"  (no hints: {e})")

        # Regenerating must never destroy hours of labelling. A filled sheet
        # is data, not a build artefact.
        if sheet.exists():
            try:
                existing = read_gold(sheet)
            except Exception:
                existing = {}
            if existing:
                sys.exit(f"{sheet} already holds {len(existing)} labels — "
                         f"refusing to overwrite.\n"
                         f"Score it:   --sheet {sheet.name} --score <taxonomy.json>\n"
                         f"Or write elsewhere:  --sheet nuovo.xlsx --make-sheet")
        out, n = make_sheet(counts, examples, hints, sheet, top=a.top)
        covered = sum(counts[p] for p in keep)
        print(f"\nWrote {out}")
        print(f"  {n} phrases to label, covering {covered}/{sum(counts.values())} "
              f"records ({covered / sum(counts.values()):.0%})")
        print("  Column D holds the real records — decide from those.")
        return

    if a.score:
        if not sheet.exists():
            sys.exit(f"{sheet} not found. Run --make-sheet first.")
        gold = read_gold(sheet)
        if not gold:
            sys.exit("No labels in the gold column yet.")
        standard = load_standard_terms()
        n_new = sum(1 for w, _ in gold.values() if w.upper().startswith(NUOVO))
        print(f"Gold set: {len(gold)} labelled phrases "
              f"({n_new} marked {NUOVO}*), "
              f"{len(standard)} standard terms")
        for path in a.score:
            p = Path(path)
            if not p.exists():
                p = work / path
            rows, stats = score(gold, phrase_to_category(p), standard)
            report(p.name, rows, stats)
        return

    ap.print_help()


if __name__ == "__main__":
    main()
